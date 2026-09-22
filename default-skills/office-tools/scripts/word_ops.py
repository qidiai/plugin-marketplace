"""Word document operations CLI tool. Usage: python word_ops.py <command> [args]"""
import sys
import json
import argparse

def cmd_read(args):
    from docx import Document
    doc = Document(args.input)
    result = {"paragraphs": [], "tables": []}
    for para in doc.paragraphs:
        if para.text.strip():
            result["paragraphs"].append({
                "text": para.text,
                "style": para.style.name if para.style else "Normal",
            })
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        result["tables"].append(rows)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for p in result["paragraphs"]:
            style = f" [{p['style']}]" if p["style"] != "Normal" else ""
            print(f"{p['text']}{style}")
        for i, table in enumerate(result["tables"]):
            print(f"\n--- Table {i+1} ---")
            for row in table:
                print(" | ".join(row))

def cmd_create(args):
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    sections = json.loads(args.content) if args.content.startswith("[") else [{"type": "paragraph", "text": args.content}]
    for sec in sections:
        t = sec.get("type", "paragraph")
        if t == "paragraph":
            p = doc.add_paragraph(sec.get("text", ""))
            if "style" in sec:
                p.style = sec["style"]
            if "bold" in sec or "size" in sec:
                for run in p.runs:
                    if sec.get("bold"):
                        run.bold = True
                    if "size" in sec:
                        run.font.size = Pt(sec["size"])
        elif t == "heading":
            doc.add_heading(sec.get("text", ""), level=sec.get("level", 1))
        elif t == "table":
            data = sec.get("data", [])
            if data:
                table = doc.add_table(rows=len(data), cols=len(data[0]))
                table.style = sec.get("table_style", "Table Grid")
                for i, row in enumerate(data):
                    for j, cell_text in enumerate(row):
                        table.rows[i].cells[j].text = str(cell_text)
        elif t == "bullet":
            doc.add_paragraph(sec.get("text", ""), style="List Bullet")
        elif t == "number":
            doc.add_paragraph(sec.get("text", ""), style="List Number")
    doc.save(args.output)
    print(f"Success: Created {args.output}")

def cmd_edit(args):
    from docx import Document
    replacements = json.loads(args.replacements)
    doc = Document(args.input)
    count = 0
    for para in doc.paragraphs:
        for old, new in replacements.items():
            if old in para.text:
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in para.text:
                            for run in para.runs:
                                if old in run.text:
                                    run.text = run.text.replace(old, new)
                                    count += 1
    doc.save(args.input if not args.output else args.output)
    print(f"Success: Replaced {count} occurrences")

def cmd_info(args):
    from docx import Document
    doc = Document(args.input)
    info = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
    }
    print(json.dumps(info, ensure_ascii=False, indent=2))

def main():
    parser = argparse.ArgumentParser(description="Word document operations tool")
    sub = parser.add_subparsers(dest="command", required=True)

    p_read = sub.add_parser("read", help="Read text from docx")
    p_read.add_argument("input", help="Input .docx file")
    p_read.add_argument("--json", action="store_true", help="Output as JSON")

    p_create = sub.add_parser("create", help="Create a new docx")
    p_create.add_argument("-o", "--output", required=True, help="Output .docx file")
    p_create.add_argument("content", help='Content: plain text or JSON array of sections. Sections: {"type":"paragraph|heading|table|bullet|number","text":"...","level":1,"data":[[]],"style":"...","bold":true,"size":12}')

    p_edit = sub.add_parser("edit", help="Replace text in docx")
    p_edit.add_argument("input", help="Input .docx file")
    p_edit.add_argument("replacements", help='JSON dict of {"old":"new"} replacements')
    p_edit.add_argument("-o", "--output", help="Output file (default: overwrite input)")

    p_info = sub.add_parser("info", help="Get document info")
    p_info.add_argument("input", help="Input .docx file")

    args = parser.parse_args()
    commands = {
        "read": cmd_read,
        "create": cmd_create,
        "edit": cmd_edit,
        "info": cmd_info,
    }
    try:
        commands[args.command](args)
    except FileNotFoundError as e:
        print(f"Error: File not found: {e.filename}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
