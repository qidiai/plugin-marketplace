"""Excel operations CLI tool. Usage: python excel_ops.py <command> [args]"""
import sys
import json
import argparse

def cmd_read(args):
    from openpyxl import load_workbook
    wb = load_workbook(args.input, data_only=args.data_only)
    sheet = wb[args.sheet] if args.sheet else wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([str(c) if c is not None else "" for c in row])
    if args.json:
        print(json.dumps({"sheet": sheet.title, "rows": rows}, ensure_ascii=False, indent=2))
    else:
        print(f"Sheet: {sheet.title} ({len(rows)} rows x {sheet.max_column} cols)")
        for row in rows[:args.limit]:
            print(" | ".join(row))
        if len(rows) > args.limit:
            print(f"... and {len(rows) - args.limit} more rows")

def cmd_create(args):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    data = json.loads(args.content)
    if isinstance(data, dict):
        data = [data]
    elif isinstance(data, list) and data and isinstance(data[0], list):
        data = [{"sheet": "Sheet1", "data": data}]
    wb = Workbook()
    wb.remove(wb.active)
    for item in data:
        sheet_name = item.get("sheet", f"Sheet{len(wb.sheetnames)+1}")
        ws = wb.create_sheet(title=sheet_name)
        table_data = item.get("data", [])
        for r_idx, row in enumerate(table_data, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        if item.get("headers") and table_data:
            for c in range(1, len(table_data[0])+1):
                cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
    wb.save(args.output)
    print(f"Success: Created {args.output} with {len(wb.sheetnames)} sheet(s)")

def cmd_edit(args):
    from openpyxl import load_workbook
    updates = json.loads(args.updates)
    wb = load_workbook(args.input)
    ws = wb[args.sheet] if args.sheet else wb.active
    count = 0
    for upd in updates:
        r, c, v = upd["row"], upd["col"], upd["value"]
        ws.cell(row=r, column=c, value=v)
        count += 1
    out = args.output or args.input
    wb.save(out)
    print(f"Success: Updated {count} cell(s) in {ws.title}")

def cmd_info(args):
    from openpyxl import load_workbook
    wb = load_workbook(args.input, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({
            "name": ws.title,
            "rows": ws.max_row,
            "cols": ws.max_column,
        })
    wb.close()
    print(json.dumps({"sheets": sheets}, ensure_ascii=False, indent=2))

def main():
    parser = argparse.ArgumentParser(description="Excel operations tool")
    sub = parser.add_subparsers(dest="command", required=True)

    p_read = sub.add_parser("read", help="Read data from xlsx")
    p_read.add_argument("input", help="Input .xlsx file")
    p_read.add_argument("--sheet", help="Sheet name (default: active sheet)")
    p_read.add_argument("--json", action="store_true", help="Output as JSON")
    p_read.add_argument("--data-only", action="store_true", help="Read computed values, not formulas")
    p_read.add_argument("--limit", type=int, default=50, help="Max rows to print (default: 50)")

    p_create = sub.add_parser("create", help="Create a new xlsx")
    p_create.add_argument("-o", "--output", required=True, help="Output .xlsx file")
    p_create.add_argument("content", help='JSON: [[row],[row]] or {"sheet":"Name","data":[[row]]} or [{"sheet":"S1","data":[...]},{"sheet":"S2","data":[...]}]')

    p_edit = sub.add_parser("edit", help="Edit cells in xlsx")
    p_edit.add_argument("input", help="Input .xlsx file")
    p_edit.add_argument("updates", help='JSON array of [{"row":1,"col":1,"value":"v"}]')
    p_edit.add_argument("--sheet", help="Sheet name (default: active sheet)")
    p_edit.add_argument("-o", "--output", help="Output file (default: overwrite input)")

    p_info = sub.add_parser("info", help="Get workbook info")
    p_info.add_argument("input", help="Input .xlsx file")

    args = parser.parse_args()
    commands = {
        "read": cmd_read,
        "create": cmd_create,
        "edit": cmd_edit,
        "info": cmd_info,
    }
    try:
        commands[args.command](args)
    except FileNotFoundError as e:
        print(f"Error: File not found: {e.filename}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
