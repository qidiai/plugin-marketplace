"""PowerPoint operations CLI tool. Usage: python ppt_ops.py <command> [args]"""
import sys
import json
import argparse

def cmd_read(args):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(args.input)
    result = []
    for i, slide in enumerate(prs.slides):
        slide_info = {"slide": i + 1, "shapes": []}
        for shape in slide.shapes:
            shape_info = {"type": str(shape.shape_type), "name": shape.name}
            if shape.has_text_frame:
                texts = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
                if texts:
                    shape_info["text"] = "\n".join(texts)
            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                shape_info["table"] = rows
            slide_info["shapes"].append(shape_info)
        result.append(slide_info)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for s in result:
            print(f"\n=== Slide {s['slide']} ===")
            for shape in s["shapes"]:
                if "text" in shape:
                    print(f"  [{shape['name']}] {shape['text']}")
                if "table" in shape:
                    print(f"  [Table: {shape['name']}]")
                    for row in shape["table"]:
                        print("    " + " | ".join(row))

def cmd_create(args):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slides_data = json.loads(args.content)
    if isinstance(slides_data, dict):
        slides_data = [slides_data]
    for sd in slides_data:
        layout_idx = sd.get("layout", 1)
        if layout_idx >= len(prs.slide_layouts):
            layout_idx = 1
        layout = prs.slide_layouts[layout_idx]
        slide = prs.slides.add_slide(layout)
        if sd.get("title"):
            if slide.shapes.title:
                slide.shapes.title.text = sd["title"]
        if sd.get("subtitle"):
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = sd["subtitle"]
        if sd.get("content"):
            left = Inches(1)
            top = Inches(2)
            width = Inches(8)
            height = Inches(4)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            txBox.text_frame.text = sd["content"]
    prs.save(args.output)
    print(f"Success: Created {args.output} with {len(slides_data)} slide(s)")

def cmd_info(args):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(args.input)
    info = {
        "slides": len(prs.slides),
        "slide_width": str(prs.slide_width),
        "slide_height": str(prs.slide_height),
        "layouts": [l.name for l in prs.slide_layouts],
    }
    print(json.dumps(info, ensure_ascii=False, indent=2))

def main():
    parser = argparse.ArgumentParser(description="PowerPoint operations tool")
    sub = parser.add_subparsers(dest="command", required=True)

    p_read = sub.add_parser("read", help="Read content from pptx")
    p_read.add_argument("input", help="Input .pptx file")
    p_read.add_argument("--json", action="store_true", help="Output as JSON")

    p_create = sub.add_parser("create", help="Create a new pptx")
    p_create.add_argument("-o", "--output", required=True, help="Output .pptx file")
    p_create.add_argument("content", help='JSON: {"title":"T","subtitle":"S","content":"C","layout":1} or array of slide objects')

    p_info = sub.add_parser("info", help="Get presentation info")
    p_info.add_argument("input", help="Input .pptx file")

    args = parser.parse_args()
    commands = {
        "read": cmd_read,
        "create": cmd_create,
        "info": cmd_info,
    }
    try:
        commands[args.command](args)
    except FileNotFoundError as e:
        print(f"Error: File not found: {e.filename}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
