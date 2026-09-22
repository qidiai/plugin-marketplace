#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
施工进度横道图生成器
在Word文档中用段落顶部边框绘制连续横道线。

用法：
  python generate_gantt.py --output output.docx --tasks "施工准备,0,10" "测量放线,5,25" --total-days 180
"""

import argparse
import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ============================================================
# 通用工具函数
# ============================================================

def set_table_all_borders(table, sz="4"):
    """设置表格所有边框（sz=4 表示 0.5pt）"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is None:
        tblBorders = OxmlElement('w:tblBorders')
        tblPr.append(tblBorders)
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        elem = tblBorders.find(qn(f'w:{edge}'))
        if elem is None:
            elem = OxmlElement(f'w:{edge}')
            tblBorders.append(elem)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), sz)
        elem.set(qn('w:space'), '0')
        elem.set(qn('w:color'), '000000')


def set_cell_diagonal(cell):
    """设置单元格对角线"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    diag = OxmlElement('w:tlToBr')
    diag.set(qn('w:val'), 'single')
    diag.set(qn('w:sz'), '4')
    diag.set(qn('w:space'), '0')
    diag.set(qn('w:color'), '000000')
    tcBorders.append(diag)


def set_cell_text(cell, text, font_size=10.5, align='center', font_name='宋体'):
    """设置单元格文字 - 默认五号宋体(10.5pt)，行距28磅"""
    cell.text = ''
    p = cell.paragraphs[0]
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            run = p.add_run()
            run.add_break()
        run = p.add_run(line)
        run.font.size = Pt(font_size)
        run.font.name = font_name
        run.font.color.rgb = None
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'left':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # 行距28磅 = 560 twips
    pPr = p._element.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:line'), '560')
    spacing.set(qn('w:lineRule'), 'exact')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_row_height(row, height_cm, rule='atLeast'):
    """设置行高"""
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_cm * 567)))
    if rule == 'exact':
        trHeight.set(qn('w:hRule'), 'exact')
    else:
        trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)


def set_row_cant_split(row):
    """设置行不可跨页拆分"""
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)


def set_col_width(cell, width_cm):
    cell.width = Cm(width_cm)


def set_cell_margins_zero(cell):
    """设置单元格左右边距为0，使段落边框无缝连接"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side in ['left', 'right']:
        elem = tcMar.find(qn(f'w:{side}'))
        if elem is None:
            elem = OxmlElement(f'w:{side}')
            tcMar.append(elem)
        elem.set(qn('w:w'), '0')
        elem.set(qn('w:type'), 'dxa')


def set_paragraph_top_border(paragraph, sz=16):
    """设置段落顶部边框（sz=16 表示 2pt）"""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    top = pBdr.find(qn('w:top'))
    if top is None:
        top = OxmlElement('w:top')
        pBdr.append(top)
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), str(sz))
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), '000000')


def set_paragraph_spacing_centered(paragraph, row_height_twips, border_sz=16):
    """设置段落间距使顶部边框在行内居中

    原理：行距设为极小值(1pt=20twips)，段前间距=(行高-边框厚度)/2
    """
    pPr = paragraph._element.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:line'), '20')
    spacing.set(qn('w:lineRule'), 'exact')
    border_twips = int(border_sz * 2.5)
    before = int((row_height_twips - border_twips) / 2)
    spacing.set(qn('w:before'), str(before))
    spacing.set(qn('w:after'), '0')


def set_paragraph_indent(paragraph, left_twips=0, right_twips=0):
    """设置段落左右缩进（用于精确控制横道线起止位置）"""
    pPr = paragraph._element.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if left_twips > 0:
        ind.set(qn('w:left'), str(left_twips))
        ind.set(qn('w:leftChars'), '0')
    if right_twips > 0:
        ind.set(qn('w:right'), str(right_twips))
        ind.set(qn('w:rightChars'), '0')


def add_chart_title(doc, text):
    """图表标题：黑色四号宋体(14pt)，两端对齐，缩进两字符，行距28磅"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(14)
    run.font.name = '宋体'
    run.font.color.rgb = None
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体')
    rFonts.set(qn('w:ascii'), '宋体')
    rFonts.set(qn('w:hAnsi'), '宋体')
    rPr.append(rFonts)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p._element.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLineChars'), '200')
    ind.set(qn('w:firstLine'), '560')
    pPr.append(ind)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:line'), '560')
    spacing.set(qn('w:lineRule'), 'exact')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    pPr.append(spacing)
    return p


def add_para_small(doc):
    """极小段落用于表间分隔"""
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:line'), '20')
    spacing.set(qn('w:lineRule'), 'exact')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    pPr.append(spacing)
    return p


# ============================================================
# 横道图核心函数
# ============================================================

def create_gantt(doc, tasks, total_days=180, title="附表四：计划开、竣工日期和施工进度网络图",
                 subtitle="扎赉诺尔区回春路、康乐路、富强路道路提升改造工程",
                 time_interval=10, name_col_width=5.0, row_height_cm=1.0,
                 border_sz="4", bar_sz=16):
    """
    在doc文档中创建横道图表格。

    参数：
        doc: Document 对象
        tasks: [(工序名称, 开始天数, 结束天数), ...]
        total_days: 总工期天数
        title: 图表标题
        subtitle: 副标题（工程名称）
        time_interval: 每列代表的天数（默认10天）
        name_col_width: 工序名称列宽度(cm)
        row_height_cm: 数据行高度(cm)
        border_sz: 表格边框粗细(默认"4"=0.5pt)
        bar_sz: 横道线粗细(默认16=2pt)
    """

    N_TIME = total_days // time_interval  # 时间列数
    N_TASKS = len(tasks)

    # 可用宽度 = 页面宽度 - 左右页边距（A4横向：29.7 - 4 = 25.7cm）
    available_width = 25.7
    time_col_width = (available_width - name_col_width) / N_TIME
    time_col_twips = int(time_col_width * 567)  # cm → twips
    row_h_twips = int(row_height_cm * 567)

    # 标题
    add_chart_title(doc, title)
    add_chart_title(doc, subtitle)
    add_chart_title(doc, f"计划工期：{total_days}日历天（以实际开工日期为准）")

    # 创建表格
    n_rows = N_TASKS + 1  # 1行表头 + N_TASKS行数据
    n_cols = N_TIME + 1   # 1列工序名 + N_TIME列时间
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False

    # 列宽
    for row in table.rows:
        set_col_width(row.cells[0], name_col_width)
        for i in range(1, n_cols):
            set_col_width(row.cells[i], time_col_width)
        set_row_cant_split(row)

    # 行高
    set_row_height(table.rows[0], 1.5, 'atLeast')  # 表头行
    for i in range(1, n_rows):
        set_row_height(table.rows[i], row_height_cm, 'atLeast')

    # 表头：斜线 + 时间刻度
    set_cell_diagonal(table.cell(0, 0))
    set_cell_text(table.cell(0, 0), "内容\n时间", font_size=10.5)
    for i in range(N_TIME):
        set_cell_text(table.cell(0, i + 1), str((i + 1) * time_interval), font_size=10.5)

    # 工序名称
    for row_idx, (name, start, end) in enumerate(tasks):
        set_cell_text(table.cell(row_idx + 1, 0), name, font_size=10.5)

    # 横道线：段落顶部边框 + 零边距 + 缩放
    for row_idx, (name, start, end) in enumerate(tasks):
        excel_row = row_idx + 1

        # 所有时间区域单元格设零边距
        for col in range(1, n_cols):
            set_cell_margins_zero(table.cell(excel_row, col))

        # 计算横道线起止列和缩放
        start_col = start // time_interval
        end_col = end // time_interval
        start_frac = (start % time_interval) / float(time_interval)
        end_frac = (end % time_interval) / float(time_interval)

        bar_start_col = start_col
        bar_end_col = end_col if end_frac > 0 else end_col - 1

        if bar_end_col < bar_start_col:
            # 横道线在同一列内
            cell = table.cell(excel_row, bar_start_col + 1)
            p = cell.paragraphs[0]
            left_indent = int(start_frac * time_col_twips)
            right_indent = int((1 - end_frac) * time_col_twips)
            set_paragraph_indent(p, left_twips=left_indent, right_twips=right_indent)
            set_paragraph_top_border(p, sz=bar_sz)
            set_paragraph_spacing_centered(p, row_h_twips, bar_sz)
        else:
            # 起始列：左缩进
            cell_start = table.cell(excel_row, bar_start_col + 1)
            p_start = cell_start.paragraphs[0]
            left_indent = int(start_frac * time_col_twips)
            set_paragraph_indent(p_start, left_twips=left_indent, right_twips=0)
            set_paragraph_top_border(p_start, sz=bar_sz)
            set_paragraph_spacing_centered(p_start, row_h_twips, bar_sz)

            # 中间列：无缩进
            for col in range(bar_start_col + 1, bar_end_col + 1):
                cell_mid = table.cell(excel_row, col + 1)
                p_mid = cell_mid.paragraphs[0]
                set_paragraph_indent(p_mid, left_twips=0, right_twips=0)
                set_paragraph_top_border(p_mid, sz=bar_sz)
                set_paragraph_spacing_centered(p_mid, row_h_twips, bar_sz)

            # 结束列：右缩进
            if end_frac > 0:
                cell_end = table.cell(excel_row, bar_end_col + 1)
                p_end = cell_end.paragraphs[0]
                right_indent = int((1 - end_frac) * time_col_twips)
                set_paragraph_indent(p_end, left_twips=0, right_twips=right_indent)
                set_paragraph_top_border(p_end, sz=bar_sz)
                set_paragraph_spacing_centered(p_end, row_h_twips, bar_sz)

    # 全表边框
    set_table_all_borders(table, sz=border_sz)
    return table


# ============================================================
# 命令行入口
# ============================================================

def main():
    parser = argparse.ArgumentParser(description='生成施工进度横道图Word文档')
    parser.add_argument('--output', '-o', required=True, help='输出文件路径')
    parser.add_argument('--tasks', '-t', nargs='+', required=True,
                        help='工序列表，格式："名称,开始天,结束天"，如 "施工准备,0,10"')
    parser.add_argument('--total-days', '-d', type=int, default=180, help='总工期天数（默认180）')
    parser.add_argument('--title', default='附表四：计划开、竣工日期和施工进度网络图', help='图表标题')
    parser.add_argument('--subtitle', default='', help='副标题（工程名称）')
    args = parser.parse_args()

    # 解析工序列表
    tasks = []
    for t in args.tasks:
        parts = t.split(',')
        if len(parts) != 3:
            print(f"错误: 工序格式应为 '名称,开始天,结束天'， got: {t}")
            return
        tasks.append((parts[0], int(parts[1]), int(parts[2])))

    # 创建文档
    doc = Document()

    # A4横向
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    create_gantt(doc, tasks, total_days=args.total_days,
                 title=args.title, subtitle=args.subtitle)

    # 确保输出目录存在
    output_dir = os.path.dirname(args.output)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)

    doc.save(args.output)
    print(f"文档已保存: {args.output}")


if __name__ == '__main__':
    main()
